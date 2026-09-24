"""Build the Aegis PS-8 slide deck (docs/aegis-ps8.pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

BG = RGBColor(0x11, 0x12, 0x17)
PANEL = RGBColor(0x18, 0x1B, 0x1F)
BORDER = RGBColor(0x2C, 0x32, 0x35)
TEXT = RGBColor(0xCC, 0xCC, 0xDC)
MUTED = RGBColor(0x8B, 0x8D, 0x97)
FAINT = RGBColor(0x6D, 0x70, 0x7A)
ORANGE = RGBColor(0xFF, 0x98, 0x30)
BLUE = RGBColor(0x57, 0x94, 0xF2)
PURPLE = RGBColor(0xB8, 0x77, 0xD9)
GREEN = RGBColor(0x73, 0xBF, 0x69)
RED = RGBColor(0xF2, 0x49, 0x5C)

FONT = "Inter"
MONO = "JetBrains Mono"
M = 0.55
CW = 13.333 - 2 * M
FOOT = 6.98

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    fill = s.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return s


def rect(s, x, y, w, h, fill=None, line=None, radius=0.05,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, paras, size=13, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.12, font=FONT):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = para.get("align", align)
        p.line_spacing = para.get("spacing", spacing)
        p.space_after = Pt(para.get("after", 0))
        p.space_before = Pt(para.get("before", 0))
        for run in para["parts"]:
            r = p.add_run()
            r.text = run["t"]
            r.font.name = run.get("font", font)
            r.font.size = Pt(run.get("size", size))
            r.font.bold = run.get("bold", bold)
            r.font.color.rgb = run.get("color", color)
    return box


def header(s, kicker, title):
    rect(s, M, 0.48, 0.3, 0.3, ORANGE, radius=0.22)
    text(s, M, 0.48, 0.3, 0.3, [{"parts": [{"t": "A", "color": BG, "bold": True, "size": 13}]}],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M + 0.44, 0.46, 8, 0.25,
         [{"parts": [{"t": kicker, "color": FAINT, "size": 10.5, "bold": True}]}])
    text(s, M, 0.74, 12, 0.6,
         [{"parts": [{"t": title, "color": TEXT, "size": 26, "bold": True}]}])


def footer(s, n):
    text(s, M, FOOT, 4, 0.25,
         [{"parts": [{"t": "Aegis · PS-8", "color": FAINT, "size": 9}]}])
    text(s, 13.333 - M - 1, FOOT, 1, 0.25,
         [{"parts": [{"t": "%d / 6" % n, "color": FAINT, "size": 9}]}],
         align=PP_ALIGN.RIGHT)


def card(s, x, y, w, h, accent, title, body):
    rect(s, x, y, 0.055, h, accent, shape=MSO_SHAPE.RECTANGLE)
    rect(s, x + 0.055, y, w - 0.055, h, PANEL, BORDER)
    text(s, x + 0.32, y + 0.3, w - 0.68, 0.8,
         [{"parts": [{"t": title, "color": TEXT, "size": 15.5, "bold": True}]}],
         spacing=1.05)
    text(s, x + 0.32, y + 1.18, w - 0.68, h - 1.45,
         [{"parts": [{"t": body, "color": MUTED, "size": 12.5}]}],
         spacing=1.25)


''' 1 · title '''
s = slide()
rect(s, 6.3665, 1.78, 0.6, 0.6, ORANGE, radius=0.24)
text(s, 6.3665, 1.78, 0.6, 0.6,
     [{"parts": [{"t": "A", "color": BG, "bold": True, "size": 22}]}],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.5, 2.62, 10.333, 0.3,
     [{"parts": [{"t": "SMART INDIA HACKATHON · PROBLEM STATEMENT 8", "color": FAINT,
                 "size": 11.5, "bold": True}]}],
     align=PP_ALIGN.CENTER)
text(s, 1.5, 2.92, 10.333, 1.0,
     [{"parts": [{"t": "Aegis", "color": TEXT, "size": 52, "bold": True}]}],
     align=PP_ALIGN.CENTER)
text(s, 1.5, 4.0, 10.333, 0.4,
     [{"parts": [{"t": "Insider threat, anomalous access and autonomous cyber defense.",
                 "color": MUTED, "size": 17}]}],
     align=PP_ALIGN.CENTER)
rect(s, 6.1665, 4.62, 1.0, 0.035, ORANGE, shape=MSO_SHAPE.RECTANGLE)
chips = ["self-hosted", "sub-millisecond guard", "one console"]
chip_w, chip_gap = 2.75, 0.22
x0 = (13.333 - (3 * chip_w + 2 * chip_gap)) / 2
for i, label in enumerate(chips):
    x = x0 + i * (chip_w + chip_gap)
    rect(s, x, 4.95, chip_w, 0.46, PANEL, BORDER, radius=0.5)
    text(s, x, 4.95, chip_w, 0.46,
         [{"parts": [{"t": label, "color": MUTED, "size": 11.5}]}],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

''' 2 · problem '''
s = slide()
header(s, "PROBLEM", "What goes wrong today")
card(s, M, 1.72, 3.911, 3.3, RED, "Attacks look like normal traffic",
     "Prompt injection hijacks the model through a chat message. Scrapers and "
     "login bots hide inside ordinary API calls, so firewalls see nothing to stop.")
card(s, M + 4.161, 1.72, 3.911, 3.3, ORANGE, "Insiders don't trip rate limits",
     "One request to a file someone has never opened is invisible to volume rules. "
     "Behaviour, not volume, gives them away: new resources, odd hours, shared keys.")
card(s, M + 8.322, 1.72, 3.911, 3.3, PURPLE, "Tools alert, they do not act",
     "Dashboards report the damage after the data or keys are gone. The block has to "
     "happen while the request is still in flight.")
rect(s, M, 5.35, CW, 0.62, PANEL, BORDER)
text(s, M, 5.35, CW, 0.62,
     [{"parts": [
         {"t": "Aegis watches behaviour, not just signatures, and ", "color": MUTED, "size": 12.5},
         {"t": "acts while the request is in flight.", "color": TEXT, "size": 12.5},
     ]}],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)

''' 3 · product '''
s = slide()
header(s, "PRODUCT", "One guard for models, APIs and insiders")
cw3 = (CW - 0.28) / 2
card(s, M, 1.68, cw3, 2.14, BLUE, "Prompt-injection firewall",
     "Screens every request before the model sees it. 22/22 on the evaluation set, "
     "and a distilled student answers in microseconds.")
card(s, M + cw3 + 0.28, 1.68, cw3, 2.14, ORANGE, "API abuse detector",
     "Builds a behavioural profile per client: rate, cadence, ID walks, auth failures. "
     "Blocks scrapers and credential stuffing automatically.")
card(s, M, 3.98, cw3, 2.14, PURPLE, "Insider access graph",
     "Learns which resources each client normally touches, then flags never-seen access, "
     "rare endpoints and keys shared across addresses.")
card(s, M + cw3 + 0.28, 3.98, cw3, 2.14, GREEN, "Console and response",
     "One dashboard for monitoring, settings and retention, with block and unblock "
     "controls and a full decision log.")
footer(s, 3)

''' 4 · how it works '''
s = slide()
header(s, "HOW IT WORKS", "Triggers first, model second")
labels = ["requests &\ntelemetry", "cheap\ntriggers", "Laya decision\nmodels",
          "allow / block\npolicy", "enforcement\n+ log"]
box_w = (CW - 4 * 0.30) / 5
box_y, box_h = 1.8, 1.35
for i, label in enumerate(labels):
    x = M + i * (box_w + 0.30)
    rect(s, x, box_y, box_w, box_h, PANEL, BORDER)
    paras = [{"parts": [{"t": line, "color": TEXT, "size": 12}]} for line in label.split("\n")]
    text(s, x + 0.1, box_y, box_w - 0.2, box_h, paras,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.15)
    if i < 4:
        ax = x + box_w + (0.30 - 0.22) / 2
        rect(s, ax, box_y + box_h / 2 - 0.09, 0.22, 0.18, FAINT,
             shape=MSO_SHAPE.RIGHT_ARROW)
bullets = [
    "Triggers keep cost low: the decision model runs only when traffic looks off.",
    "Every verdict carries evidence: thresholds, probabilities and the narrative the model read.",
    "Runs as one self-hosted Python process with SQLite; no cloud call per decision.",
    "The console reads the same log: every block, unblock and setting change is visible live.",
]
paras = []
for b in bullets:
    paras.append({"parts": [
        {"t": "•  ", "color": ORANGE, "bold": True, "size": 14},
        {"t": b, "color": TEXT, "size": 14},
    ], "after": 14, "spacing": 1.2})
text(s, M, 3.75, CW, 2.4, paras)
footer(s, 4)

''' 5 · evidence '''
s = slide()
header(s, "EVIDENCE", "Measured on this machine")
stats = [
    (GREEN, "22/22", "prompt-injection evaluation",
     "rules plus model, with no false blocks on legitimate security work"),
    (BLUE, "6/6", "client profiles classified",
     "normal, legit poller, scraper, credential stuffing, burst, fuzzer"),
    (RED, "~11 s", "scraper caught mid-run",
     "900+ following requests rejected with 429 until the block expired"),
    (ORANGE, "309 µs", "median fast-path decision",
     "about 2 ms end to end, versus about 0.9 s for the full model"),
]
sw = (CW - 0.28) / 2
sh_ = 1.62
positions = [(M, 1.72), (M + sw + 0.28, 1.72), (M, 3.56), (M + sw + 0.28, 3.56)]
for (x, y), (color, big, label, sub) in zip(positions, stats):
    rect(s, x, y, 0.055, sh_, color, shape=MSO_SHAPE.RECTANGLE)
    rect(s, x + 0.055, y, sw - 0.055, sh_, PANEL, BORDER)
    text(s, x + 0.35, y + 0.22, sw - 0.7, 0.65,
         [{"parts": [{"t": big, "color": color, "size": 32, "bold": True}]}])
    text(s, x + 0.35, y + 0.88, sw - 0.7, 0.3,
         [{"parts": [{"t": label, "color": TEXT, "size": 13.5, "bold": True}]}])
    text(s, x + 0.35, y + 1.17, sw - 0.7, 0.4,
         [{"parts": [{"t": sub, "color": MUTED, "size": 11.5}]}], spacing=1.1)
text(s, M, 5.5, CW, 0.3,
     [{"parts": [{"t": "All figures measured on this machine and reproducible through the "
                       "benches in guard/tools/.",
                 "color": FAINT, "size": 11}]}],
     align=PP_ALIGN.CENTER)
footer(s, 5)

''' 6 · status '''
s = slide()
header(s, "STATUS", "Where it stands today")
cw6 = (CW - 0.28) / 2
items_left = [
    "End-to-end enforcement: injection and abuse blocks applied while the request is in flight.",
    "Console: first-run setup, sessions, live metrics, settings and retention.",
    "Demos: chat firewall and abuse bench sending real requests.",
]
items_right = [
    "Wire graph features into the live decision narrative.",
    "Retrain the fast path on real customer traffic.",
    "Add roles to the console and export alerts to a SIEM.",
]


def list_card(x, y, w, h, accent, title, items):
    rect(s, x, y, 0.055, h, accent, shape=MSO_SHAPE.RECTANGLE)
    rect(s, x + 0.055, y, w - 0.055, h, PANEL, BORDER)
    text(s, x + 0.32, y + 0.28, w - 0.64, 0.35,
         [{"parts": [{"t": title, "color": TEXT, "size": 15.5, "bold": True}]}])
    paras = []
    for item in items:
        paras.append({"parts": [
            {"t": "•  ", "color": accent, "bold": True, "size": 13},
            {"t": item, "color": MUTED, "size": 13},
        ], "after": 10, "spacing": 1.2})
    text(s, x + 0.32, y + 0.82, w - 0.64, h - 1.05, paras)


list_card(M, 1.72, cw6, 3.3, GREEN, "Working today", items_left)
list_card(M + cw6 + 0.28, 1.72, cw6, 3.3, BLUE, "Next", items_right)
rect(s, M, 5.35, CW, 0.62, PANEL, BORDER)
text(s, M, 5.35, CW, 0.62,
     [{"parts": [
         {"t": "Run it:  ", "color": FAINT, "size": 12.5, "bold": True, "font": MONO},
         {"t": ".venv/bin/python guard/app.py", "color": TEXT, "size": 12.5, "font": MONO},
         {"t": "   →   ", "color": FAINT, "size": 12.5, "font": MONO},
         {"t": "http://127.0.0.1:8978/console", "color": ORANGE, "size": 12.5, "font": MONO},
     ]}],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 6)

prs.save("docs/aegis-ps8.pptx")
print("saved docs/aegis-ps8.pptx with", len(prs.slides._sldIdLst), "slides")
